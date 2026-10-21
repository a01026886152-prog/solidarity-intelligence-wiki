"""의정지원 Markdown 문서를 DOCX로 바꾼다 (한글 프로그램에서 열어 HWP로 저장 가능).

사용법:
    python3 uijeong/scripts/md2docx.py <입력.md> [출력.docx]

지원 문법: 제목(#~####), 문단, 글머리표(-, *), 번호 목록(1.), 체크박스(- [ ]),
인용(>), 표(| a | b |), 구분선(---), 굵게(**...**).
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "맑은 고딕"
BODY_PT = 10.5


def _set_font(run, size=None, bold=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _add_inline(paragraph, text, size=BODY_PT):
    # **굵게** 구간을 나눠 런으로 넣는다. [[링크]]는 괄호만 걷어낸다.
    text = re.sub(r"\[\[([^\]|]+)(?:\|([^\]]+))?\]\]", lambda m: m.group(2) or m.group(1), text)
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1(\2)", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            _set_font(paragraph.add_run(part), size, bold=(i % 2 == 1))


def _base_style(document):
    style = document.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(BODY_PT)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    for level in range(1, 5):
        heading = document.styles[f"Heading {level}"]
        heading.font.name = FONT
        heading.font.color.rgb = RGBColor(0, 0, 0)
        heading.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def _flush_table(document, rows):
    rows = [r for r in rows if not re.fullmatch(r"\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?", r.strip())]
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    width = max(len(r) for r in cells)
    table = document.add_table(rows=len(cells), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        for j in range(width):
            cell = table.cell(i, j)
            cell.paragraphs[0].text = ""
            _add_inline(cell.paragraphs[0], row[j] if j < len(row) else "", size=9.5)
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    document.add_paragraph()


def convert(md_path, docx_path):
    document = Document()
    _base_style(document)
    table_rows = []
    for raw in Path(md_path).read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if line.lstrip().startswith("|"):
            table_rows.append(line)
            continue
        if table_rows:
            _flush_table(document, table_rows)
            table_rows = []
        stripped = line.strip()
        if not stripped or stripped.startswith("<!--"):
            continue
        if m := re.match(r"^(#{1,4})\s+(.*)", stripped):
            level = len(m.group(1))
            p = document.add_heading(level=level)
            _add_inline(p, m.group(2), size={1: 16, 2: 13.5, 3: 12, 4: 11}[level])
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif re.fullmatch(r"-{3,}|\*{3,}", stripped):
            document.add_paragraph("─" * 40).alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif m := re.match(r"^(\s*)[-*]\s+\[( |x|X)\]\s+(.*)", line):
            p = document.add_paragraph(style="List Bullet")
            _add_inline(p, ("☑ " if m.group(2).strip() else "☐ ") + m.group(3))
        elif m := re.match(r"^(\s*)[-*]\s+(.*)", line):
            style = "List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet"
            _add_inline(document.add_paragraph(style=style), m.group(2))
        elif m := re.match(r"^(\s*)(\d+)[.)]\s+(.*)", line):
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Pt(12 + 12 * (len(m.group(1)) // 2))
            _add_inline(p, f"{m.group(2)}. {m.group(3)}")
        elif stripped.startswith(">"):
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _add_inline(p, stripped.lstrip("> ").strip())
            for run in p.runs:
                run.italic = True
        else:
            _add_inline(document.add_paragraph(), stripped)
    if table_rows:
        _flush_table(document, table_rows)
    document.save(docx_path)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = Path(sys.argv[1])
    dest = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    convert(src, dest)
    print(dest)


if __name__ == "__main__":
    main()
