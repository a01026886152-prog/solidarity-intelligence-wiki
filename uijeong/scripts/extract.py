"""의정지원 자료 텍스트 추출기.

HWP, HWPX, PDF, DOCX, TXT, MD 파일에서 본문 텍스트를 뽑아 .txt로 저장한다.

사용법:
    python3 -I uijeong/scripts/extract.py <파일 또는 폴더> [...] --out <출력폴더>

출력 폴더에 원본파일명.txt가 생기고, 파일별 글자 수를 표준출력으로 알려준다.
"""
import argparse
import re
import struct
import subprocess
import sys
import zipfile
import zlib
from pathlib import Path
from xml.etree import ElementTree

SUPPORTED = {".hwp", ".hwpx", ".pdf", ".docx", ".txt", ".md"}

# HWP 5.0 PARA_TEXT 안의 제어문자 중 16바이트(8 WCHAR)를 차지하는 확장/인라인 제어
_WIDE_CTRL = {1, 2, 3, 4, 5, 6, 7, 8, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23}
_HWPTAG_PARA_TEXT = 67


def _hwp_records(data):
    i = 0
    while i + 4 <= len(data):
        header = struct.unpack("<I", data[i:i + 4])[0]
        i += 4
        tag, size = header & 0x3FF, (header >> 20) & 0xFFF
        if size == 0xFFF:
            size = struct.unpack("<I", data[i:i + 4])[0]
            i += 4
        yield tag, data[i:i + size]
        i += size


def _hwp_para_text(body):
    chars, j = [], 0
    while j + 2 <= len(body):
        c = struct.unpack("<H", body[j:j + 2])[0]
        if c in _WIDE_CTRL:
            j += 16
        elif c == 9:
            chars.append("\t")
            j += 16
        elif c in (10, 13):
            chars.append("\n")
            j += 2
        elif c < 32:
            j += 2
        else:
            chars.append(chr(c))
            j += 2
    return "".join(chars).rstrip("\n")


def extract_hwp(path):
    import olefile

    ole = olefile.OleFileIO(str(path))
    header = ole.openstream("FileHeader").read()
    flags = struct.unpack("<I", header[36:40])[0]
    if flags & 0x04:
        raise ValueError("배포용(암호화) HWP 문서입니다. 한글에서 PDF로 저장해 다시 올려주세요.")
    compressed = flags & 0x01
    sections = sorted(
        (e for e in ole.listdir() if e[0] == "BodyText" and e[1].startswith("Section")),
        key=lambda e: int(e[1][7:]),
    )
    paragraphs = []
    for entry in sections:
        data = ole.openstream(entry).read()
        if compressed:
            data = zlib.decompress(data, -15)
        for tag, body in _hwp_records(data):
            if tag == _HWPTAG_PARA_TEXT:
                paragraphs.append(_hwp_para_text(body))
    return "\n".join(paragraphs)


def extract_hwpx(path):
    paragraphs = []
    with zipfile.ZipFile(path) as z:
        names = sorted(
            (n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n)),
            key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)),
        )
        for name in names:
            root = ElementTree.fromstring(z.read(name))
            for el in root.iter():
                if el.tag.endswith("}p"):
                    text = "".join(t.text or "" for t in el.iter() if t.tag.endswith("}t"))
                    if text:
                        paragraphs.append(text)
    return "\n".join(paragraphs)


def extract_pdf(path):
    try:
        out = subprocess.run(
            ["pdftotext", "-layout", str(path), "-"], capture_output=True, check=True
        )
        return out.stdout.decode("utf-8", "replace")
    except FileNotFoundError:
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def extract_docx(path):
    import docx

    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


EXTRACTORS = {
    ".hwp": extract_hwp,
    ".hwpx": extract_hwpx,
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".txt": lambda p: Path(p).read_text(encoding="utf-8", errors="replace"),
    ".md": lambda p: Path(p).read_text(encoding="utf-8", errors="replace"),
}


def iter_inputs(paths):
    for p in map(Path, paths):
        if p.is_dir():
            yield from sorted(f for f in p.rglob("*") if f.suffix.lower() in SUPPORTED)
        elif p.suffix.lower() in SUPPORTED:
            yield p
        else:
            print(f"건너뜀(지원하지 않는 형식): {p}", file=sys.stderr)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("inputs", nargs="+")
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    failed = 0
    for src in iter_inputs(args.inputs):
        try:
            text = EXTRACTORS[src.suffix.lower()](src)
        except Exception as e:  # 한 파일이 실패해도 나머지는 계속 처리한다
            print(f"실패: {src} — {e}", file=sys.stderr)
            failed += 1
            continue
        # HWP에 섞인 짝 없는 서로게이트 문자를 정리한다
        text = text.encode("utf-16", "surrogatepass").decode("utf-16", "replace")
        dest = out_dir / (src.stem + ".txt")
        dest.write_text(text, encoding="utf-8")
        print(f"{len(text):>8,}자  {src.name} → {dest}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
